from __future__ import annotations

import csv
import json
import shutil
from dataclasses import asdict, dataclass
from pathlib import Path

import numpy as np

import configuration as cfg
from class_mapping import canonical_class_name, class_index_map
from image_utils import list_image_files


@dataclass
class PredictionRecord:
    image_path: str
    true_index: int
    true_class: str
    pred_index: int
    pred_class: str
    score: float
    correct: bool


def collect_evaluation_samples(dataset_dir: Path, class_names: list[str], smoke_max_per_class: int = 0):
    mapping = class_index_map(class_names)
    rows = []
    for source_folder in sorted(p for p in Path(dataset_dir).iterdir() if p.is_dir()):
        class_name = canonical_class_name(source_folder.name)
        if class_name not in mapping:
            print(f"[skip] class folder not in class_names.json: {source_folder}")
            continue
        paths = list_image_files(source_folder)
        if smoke_max_per_class > 0:
            paths = paths[:smoke_max_per_class]
        rows.extend((path, mapping[class_name], class_name) for path in paths)
    return rows


def _matrix(records: list[PredictionRecord], num_classes: int) -> np.ndarray:
    matrix = np.zeros((num_classes, num_classes), dtype=np.int64)
    for row in records:
        matrix[row.true_index, row.pred_index] += 1
    return matrix


def _write_csv(path: Path, rows, fieldnames):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def _write_excel(path: Path, matrix: np.ndarray, class_names: list[str]):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "confusion_matrix"
    ws.cell(1, 1, "true\\pred")
    for idx, name in enumerate(class_names, start=2):
        ws.cell(1, idx, name)
        ws.cell(idx, 1, name)
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    for r in range(matrix.shape[0]):
        for c in range(matrix.shape[1]):
            cell = ws.cell(r + 2, c + 2, int(matrix[r, c]))
            if r == c and matrix[r, c] > 0:
                cell.fill = green
            elif r != c and matrix[r, c] > 0:
                cell.fill = red
            cell.alignment = Alignment(horizontal="center")
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for cell in ws["A"]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "B2"
    wb.save(path)


def finalize_report(records: list[PredictionRecord], class_names: list[str], output_dir: Path, copy_wrong: bool = True):
    output_dir = Path(output_dir)
    if output_dir.exists():
        shutil.rmtree(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    fieldnames = ["image_path", "true_index", "true_class", "pred_index", "pred_class", "score", "correct"]
    _write_csv(output_dir / "predictions.csv", [asdict(x) for x in records], fieldnames)

    matrix = _matrix(records, len(class_names))
    np.savetxt(output_dir / "confusion_matrix.csv", matrix, delimiter=",", fmt="%d")
    labeled_rows = []
    for idx, name in enumerate(class_names):
        row = {"true\\pred": name}
        row.update({class_names[c]: int(matrix[idx, c]) for c in range(len(class_names))})
        labeled_rows.append(row)
    _write_csv(output_dir / "confusion_matrix_labeled.csv", labeled_rows, ["true\\pred", *class_names])

    with (output_dir / "confusion_matrix_labeled.txt").open("w", encoding="utf-8") as handle:
        handle.write("true\\pred\t" + "\t".join(class_names) + "\n")
        for idx, name in enumerate(class_names):
            handle.write(name + "\t" + "\t".join(str(int(x)) for x in matrix[idx]) + "\n")
    _write_excel(output_dir / "confusion_matrix.xlsx", matrix, class_names)

    total = len(records)
    correct = sum(1 for x in records if x.correct)
    overall = correct / total if total else 0.0
    per_class = []
    for idx, name in enumerate(class_names):
        class_total = int(matrix[idx].sum())
        class_correct = int(matrix[idx, idx])
        per_class.append({
            "class_index": idx,
            "class_name": name,
            "total": class_total,
            "correct": class_correct,
            "accuracy": class_correct / class_total if class_total else 0.0,
        })
    _write_csv(output_dir / "per_class_accuracy.csv", per_class,
               ["class_index", "class_name", "total", "correct", "accuracy"])

    summary = {
        "total": total,
        "correct": correct,
        "wrong": total - correct,
        "accuracy": overall,
        "class_names": class_names,
    }
    (output_dir / "summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    with (output_dir / "confusion_report.txt").open("w", encoding="utf-8") as handle:
        handle.write(f"Total: {total}\nCorrect: {correct}\nWrong: {total - correct}\nAccuracy: {overall:.8f}\n\n")
        for row in per_class:
            handle.write(
                f"[{row['class_index']:02d}] {row['class_name']}: "
                f"{row['correct']}/{row['total']} = {row['accuracy']:.8f}\n"
            )

    if copy_wrong:
        wrong_root = output_dir / "wrongly_classified_images"
        for row in records:
            if row.correct:
                continue
            source = Path(row.image_path)
            safe_score = f"{row.score:.{cfg.STEP4_SCORE_DECIMALS}f}"
            destination = wrong_root / row.true_class / f"pred_{row.pred_class}" / f"{source.stem}_score_{safe_score}{source.suffix.lower()}"
            destination.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, destination)
    return summary
